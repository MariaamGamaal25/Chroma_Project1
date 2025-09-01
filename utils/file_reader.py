# # utils/file_reader.py
# import os
# import re
# import uuid
# from typing import List, Dict

# from PyPDF2 import PdfReader
# from docx import Document as DocxDocument
# from pdf2image import convert_from_path
# import pytesseract

# # ---------- Text cleaning ----------
# def clean_text(text: str) -> str:
#     """
#     Normalize OCR/text-layer output:
#     - Remove page markers / newlines spam
#     - Collapse multiple spaces
#     - Strip odd Unicode control chars
#     """
#     text = re.sub(r"\s+", " ", text)          # collapse whitespace
#     text = text.replace("ـ", "")              # remove tatweel
#     text = text.replace("\ufeff", "")         # remove BOM if present
#     text = text.replace("➢", "")             
#     text = text.strip()
#     return text

# # ---------- Readers ----------
# def _read_txt(file_path: str) -> str:
#     with open(file_path, "r", encoding="utf-8") as f:
#         return clean_text(f.read())

# def _read_docx(file_path: str) -> str:
#     doc = DocxDocument(file_path)
#     parts = []

#     for p in doc.paragraphs:
#         if p.text:
#             parts.append(p.text)

#     for table in doc.tables:
#         for row in table.rows:
#             for cell in row.cells:
#                 if cell.text:
#                     parts.append(cell.text)

#     return clean_text("\n".join(parts))

# def _read_pdf_with_text_layer(file_path: str) -> str:
#     try:
#         reader = PdfReader(file_path)
#         pages_txt = []
#         for page in reader.pages:
#             t = page.extract_text() or ""
#             if len(t.strip()) >= 10:
#                 pages_txt.append(t)
#         return clean_text("\n".join(pages_txt))
#     except Exception:
#         return ""

# def _read_pdf_with_ocr(file_path: str, ocr_lang: str = "ara+eng") -> str:
#     pages = convert_from_path(file_path)
#     out = []
#     for img in pages:
#         page_txt = pytesseract.image_to_string(img, lang=ocr_lang)
#         if page_txt.strip():
#             out.append(page_txt)
#     return clean_text("\n".join(out))

# def read_all_text(file_path: str) -> str:
#     ext = os.path.splitext(file_path)[1].lower()
#     if ext == ".txt":
#         return _read_txt(file_path)
#     if ext == ".docx":
#         return _read_docx(file_path)
#     if ext == ".pdf":
#         txt = _read_pdf_with_text_layer(file_path)
#         if txt.strip():
#             return txt
#         return _read_pdf_with_ocr(file_path)
#     raise ValueError(f"Unsupported file type: {ext}")

# # ---------- Chunking ----------
# # _HEADLINE_RE = re.compile(r"^\s*([A-Z][A-Za-z0-9 :/\-\(\)\[\]]{0,80})\s*$")

# _HEADLINE_RE = re.compile(r"^\s*([A-Za-z\u0600-\u06FF][A-Za-z0-9\u0600-\u06FF :/\-\(\)\[\]]{0,80})\s*$")



# def chunk_text(text: str, filename: str, max_len: int = 1200) -> List[Dict]:
#     lines = [ln.strip() for ln in text.splitlines()]
#     chunks: List[Dict] = []

#     current_title = None
#     buf: list[str] = []

#     def _flush():
#         nonlocal buf, current_title
#         if not buf:
#             return
#         merged = " ".join(buf).strip()
#         if not merged:
#             buf = []
#             return
#         start = 0
#         while start < len(merged):
#             piece = merged[start:start+max_len]
#             start += max_len
#             payload = f"{current_title}: {piece}" if current_title else piece
#             chunks.append({
#                 "id": str(uuid.uuid4()),
#                 "text": payload,
#                 "metadata": {
#                     "filename": filename,
#                     "section": current_title or "Paragraph"
#                 }
#             })
#         buf = []

#     has_headlines = False
#     for ln in lines:
#         if not ln:
#             continue
#         if _HEADLINE_RE.match(ln) and len(ln.split()) <= 10:
#             has_headlines = True
#             _flush()
#             current_title = ln
#         else:
#             buf.append(ln)

#     _flush()

#     if chunks:
#         return chunks

#     paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
#     if not paragraphs:
#         paragraphs = [text]

#     for i, p in enumerate(paragraphs, 1):
#         start = 0
#         while start < len(p):
#             piece = p[start:start+max_len]
#             start += max_len
#             chunks.append({
#                 "id": str(uuid.uuid4()),
#                 "text": piece,
#                 "metadata": {
#                     "filename": filename,
#                     "section": f"Paragraph {i}"
#                 }
#             })

#     return chunks

# def read_and_chunk_file(file_path: str) -> List[Dict]:
#     text = read_all_text(file_path)
#     return chunk_text(text, filename=os.path.basename(file_path))

import os
import re
import uuid
from typing import List, Dict

from PyPDF2 import PdfReader
from docx import Document as DocxDocument
from pdf2image import convert_from_path
import pytesseract
import pdfplumber

# ---------- Text cleaning ----------
def clean_text(text: str) -> str:
    text = re.sub(r"\s+", " ", text)          # collapse whitespace
    text = text.replace("ـ", "")              # remove tatweel
    text = text.replace("\ufeff", "")         # remove BOM if present
    return text.strip()

# ---------- Readers ----------
def _read_txt(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8") as f:
        return clean_text(f.read())

def _read_docx(file_path: str) -> str:
    doc = DocxDocument(file_path)
    parts = []
    for p in doc.paragraphs:
        if p.text:
            parts.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
    return clean_text("\n".join(parts))

def _read_pdf_with_text_layer(file_path: str) -> str:
    try:
        reader = PdfReader(file_path)
        pages_txt = []
        for page in reader.pages:
            t = page.extract_text() or ""
            if len(t.strip()) >= 10:
                pages_txt.append(t)
        return clean_text("\n".join(pages_txt))
    except Exception:
        return ""

def _read_pdf_with_ocr(file_path: str, ocr_lang: str = "ara+eng") -> str:
    pages = convert_from_path(file_path)
    out = []
    for img in pages:
        page_txt = pytesseract.image_to_string(img, lang=ocr_lang)
        if page_txt.strip():
            out.append(page_txt)
    return clean_text("\n".join(out))

def read_all_text(file_path: str) -> str:
    """Generic file reader for .txt, .docx, .pdf (fallback to OCR)."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".txt":
        return _read_txt(file_path)
    if ext == ".docx":
        return _read_docx(file_path)
    if ext == ".pdf":
        txt = _read_pdf_with_text_layer(file_path)
        if txt.strip():
            return txt
        return _read_pdf_with_ocr(file_path)
    raise ValueError(f"Unsupported file type: {ext}")


import os
import uuid
import pdfplumber
from typing import List, Dict

# ---------- PDF Service Block Extractor (with persistent header detection) ----------
def extract_pdf_service_blocks(file_path: str) -> List[Dict]:
    """
    Extract service-level blocks from PDF tables.
    Dynamically detects header row (must start with 'Service') and reuses it
    across all pages so headers like 'Not Covered, Covered' are not mistaken as headers.
    Returns a flat list of dicts for Chroma insertion.
    """
    docs: List[Dict] = []
    last_valid_header = None  # keep last valid header across pages

    with pdfplumber.open(file_path) as pdf:
        for page_idx, page in enumerate(pdf.pages, start=1):
            tables = page.extract_tables()
            for table in tables:
                if not table:
                    continue

                # --- Detect header row ---
                raw_header = [(cell.strip() if cell else "") for cell in table[0]]

                if raw_header and raw_header[0].lower() == "service":
                    header = raw_header
                    last_valid_header = header  # update the valid header
                    data_rows = table[1:]  # skip header
                else:
                    # Reuse last valid header if available
                    header = last_valid_header if last_valid_header else [
                        f"col_{i}" for i in range(len(raw_header))
                    ]
                    data_rows = table  # treat entire table as data

                current_service = None
                service_rows = []

                for row in data_rows:
                    if not row:
                        continue
                    cleaned_row = [(cell.strip() if cell else "") for cell in row]
                    service = cleaned_row[0] if len(cleaned_row) > 0 else None

                    # --- Flush service if a new one starts ---
                    if service and service.strip():
                        if current_service and service_rows:
                            summary_text = f"Service: {current_service} | Rows: {len(service_rows)}"
                            docs.append({
                                "id": str(uuid.uuid4()),
                                "text": summary_text,
                                "metadata": {
                                    "service": current_service,
                                    "doc_type": "service",
                                    "page": page_idx,
                                    "filename": os.path.basename(file_path),
                                    "source": os.path.basename(file_path),
                                    "headers": ",".join(header)  # store as string
                                }
                            })
                            docs.extend(service_rows)

                        current_service = service.strip()
                        service_rows = []

                    # --- Row doc with dynamic headers ---
                    row_meta = {
                        "service": current_service or service or "",
                        "doc_type": "row",
                        "page": page_idx,
                        "filename": os.path.basename(file_path),
                        "source": os.path.basename(file_path),
                        "headers": ",".join(header)  # store as string
                    }
                    for idx, val in enumerate(cleaned_row):
                        row_meta[header[idx]] = val

                    row_doc = {
                        "id": str(uuid.uuid4()),
                        "text": f"Service: {current_service or service} | " +
                                " | ".join([f"{header[i]}: {v}" for i, v in enumerate(cleaned_row)]),
                        "metadata": row_meta,
                    }
                    service_rows.append(row_doc)

                # --- Flush last service at end of table ---
                if current_service and service_rows:
                    summary_text = f"Service: {current_service} | Rows: {len(service_rows)}"
                    docs.append({
                        "id": str(uuid.uuid4()),
                        "text": summary_text,
                        "metadata": {
                            "service": current_service,
                            "doc_type": "service",
                            "page": page_idx,
                            "filename": os.path.basename(file_path),
                            "source": os.path.basename(file_path),
                            "headers": ",".join(header)
                        }
                    })
                    docs.extend(service_rows)

    return docs

# ---------- DOCX Service Block Extractor (with header detection) ----------

import os
import uuid
from typing import List, Dict
from docx import Document

# ---------- DOCX Service Block Extractor ----------
def extract_docx_service_blocks(file_path: str) -> List[Dict]:
    """
    Extract service-level blocks from Word (DOCX) tables.
    Each service = {summary doc + row docs}.
    Returns a flat list of dicts for Chroma insertion.
    """
    docs: List[Dict] = []
    doc = Document(file_path)

    for table_idx, table in enumerate(doc.tables, start=1):
        if not table or not table.rows:
            continue

        # Extract headers (first row)
        headers = [cell.text.strip() for cell in table.rows[0].cells]
        headers_line = ",".join(headers)

        current_service = None
        service_rows = []

        # Iterate over data rows
        for row in table.rows[1:]:
            cells = [cell.text.strip() for cell in row.cells]
            if not any(cells):
                continue  # skip empty row

            # Assume first column is always "Service"
            service = cells[0]

            # New service begins when column 0 has text
            if service:
                if current_service and service_rows:
                    # Flush previous service
                    summary_text = f"Service: {current_service} | Rows: {len(service_rows)}"
                    docs.append({
                        "id": str(uuid.uuid4()),
                        "text": summary_text,
                        "metadata": {
                            "service": current_service,
                            "doc_type": "service",
                            "page": table_idx,  # treat table number as "page"
                            "filename": os.path.basename(file_path),
                            "source": os.path.basename(file_path),
                            "headers": headers_line
                        }
                    })
                    docs.extend(service_rows)

                # Reset service
                current_service = service
                service_rows = []

            # Create row document with column mapping
            row_metadata = {
                headers[i]: cells[i] if i < len(cells) else "" for i in range(len(headers))
            }
            row_doc = {
                "id": str(uuid.uuid4()),
                "text": f"Service: {current_service or service} | " + " | ".join(
                    [f"{headers[i]}: {cells[i] if i < len(cells) else ''}" for i in range(len(headers))]
                ),
                "metadata": {
                    "service": current_service or service,
                    "doc_type": "row",
                    "page": table_idx,
                    "filename": os.path.basename(file_path),
                    "source": os.path.basename(file_path),
                    "headers": headers_line,
                    **row_metadata
                }
            }
            service_rows.append(row_doc)

        # Flush last service of table
        if current_service and service_rows:
            summary_text = f"Service: {current_service} | Rows: {len(service_rows)}"
            docs.append({
                "id": str(uuid.uuid4()),
                "text": summary_text,
                "metadata": {
                    "service": current_service,
                    "doc_type": "service",
                    "page": table_idx,
                    "filename": os.path.basename(file_path),
                    "source": os.path.basename(file_path),
                    "headers": headers_line
                }
            })
            docs.extend(service_rows)

    return docs
